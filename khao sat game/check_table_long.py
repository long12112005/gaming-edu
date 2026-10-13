import docx

doc = docx.Document(r'c:\Users\LONG\gaming-edu\khao sat game\Bao_cao_tien_do_Tuan_1_Nguyen_Nhat_Long.docx')

with open('table_check_long.txt', 'w', encoding='utf-8') as f:
    f.write("Tables:\n")
    for t_idx, table in enumerate(doc.tables):
        f.write(f"Table {t_idx}:\n")
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                if "Nguyễn Nhật Long" in cell.text or "Chèn ảnh" in cell.text:
                    f.write(f"  Row {r_idx}, Cell {c_idx}: {cell.text.strip()}\n")
                    for p_idx, cp in enumerate(cell.paragraphs):
                        f.write(f"    Cell Para {p_idx}: {cp.text.strip()}\n")
