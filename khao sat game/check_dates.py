import docx
import glob

files = glob.glob(r'c:\Users\LONG\gaming-edu\khao sat game\Bao_cao_tien_do_Tuan_*.docx')
files.sort()

with open('dates_check.txt', 'w', encoding='utf-8') as out:
    for f in files:
        out.write(f"{f}\n")
        try:
            doc = docx.Document(f)
            for p in doc.paragraphs:
                if "Thời gian:" in p.text:
                    out.write(f"  Time: {p.text}\n")
            
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            if "Đồng Nai, Ngày" in p.text:
                                out.write(f"  Date: {p.text}\n")
        except Exception as e:
            out.write(f"  Error: {e}\n")
