import docx
from docx.shared import Inches
import glob
import os

files = glob.glob(r'c:\Users\LONG\gaming-edu\khao sat game\Bao_cao_tien_do_Tuan_*.docx')
img_path = r'c:\Users\LONG\gaming-edu\khao sat game\chu ky Long.png'

if not os.path.exists(img_path):
    print("Image not found:", img_path)
    exit(1)

for file in files:
    try:
        doc = docx.Document(file)
        
        # Access the signature cell for student
        # In Table 0, Row 0, Cell 1
        table = doc.tables[0]
        cell = table.cell(0, 1)
        
        # Verify it's the right cell
        if "Nguyễn Nhật Long" in cell.text:
            name_idx = -1
            for i, p in enumerate(cell.paragraphs):
                if "Nguyễn Nhật Long" in p.text:
                    name_idx = i
                    break
            
            if name_idx > 0:
                # Clear all paragraphs between 0 and name_idx
                # This will remove the previously added image
                for i in range(1, name_idx):
                    cell.paragraphs[i].clear()
                
                # Insert the new image into paragraph 1
                para_to_insert = cell.paragraphs[1]
                run = para_to_insert.add_run()
                run.add_picture(img_path, width=Inches(1.5)) # 1.5 inches wide
                
                doc.save(file)
                print(f"Updated Long's signature in {file}")
            else:
                print(f"Could not find name paragraph in {file}")
        else:
            print(f"Could not find Nguyễn Nhật Long in cell in {file}")
    except Exception as e:
        print(f"Error processing {file}: {e}")
