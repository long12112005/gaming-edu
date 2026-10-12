import docx
from docx.shared import Inches
import glob
import os

files = glob.glob(r'c:\Users\LONG\gaming-edu\khao sat game\Bao_cao_tien_do_Tuan_*.docx')
img_path = r'c:\Users\LONG\gaming-edu\khao sat game\chu ky thay.jpg'

if not os.path.exists(img_path):
    print("Image not found:", img_path)
    exit(1)

for file in files:
    try:
        doc = docx.Document(file)
        
        # Access the signature cell
        # In Table 0, Row 0, Cell 0
        table = doc.tables[0]
        cell = table.cell(0, 0)
        
        # Verify it's the right cell
        if "Xác nhận của GVHD" in cell.text:
            # We want to insert the image in the empty paragraph before the name.
            # Usually Para 0 is "Xác nhận của GVHD", Para 1 is empty.
            para_to_insert = cell.paragraphs[1]
            
            # Check if there is already an image in this paragraph (maybe from previous run)
            # Actually just clear the paragraph and add the image
            para_to_insert.clear()
            run = para_to_insert.add_run()
            run.add_picture(img_path, width=Inches(1.5)) # 1.5 inches wide is a good size for signature
            
            doc.save(file)
            print(f"Added signature to {file}")
        else:
            print(f"Could not find signature cell in {file}")
    except Exception as e:
        print(f"Error processing {file}: {e}")
