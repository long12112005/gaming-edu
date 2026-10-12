import docx

doc = docx.Document(r'c:\Users\LONG\gaming-edu\khao sat game\Bao_cao_tien_do_Tuan_1_Nguyen_Nhat_Long.docx')

with open('table_check.txt', 'w', encoding='utf-8') as f:
    f.write("Paragraphs:\n")
    for i, p in enumerate(doc.paragraphs):
        if "Phan Thiện Phước" in p.text or "GVHD" in p.text:
            f.write(f"Para {i}: {p.text}\n")
            
    f.write("\nTables:\n")
    for t_idx, table in enumerate(doc.tables):
        f.write(f"Table {t_idx}:\n")
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                if "Phan Thiện Phước" in cell.text or "GVHD" in cell.text:
                    f.write(f"  Row {r_idx}, Cell {c_idx}: {cell.text.strip()}\n")
                    # Also print paragraphs inside the cell to see where to insert
                    for p_idx, cp in enumerate(cell.paragraphs):
                        f.write(f"    Cell Para {p_idx}: {cp.text.strip()}\n")
