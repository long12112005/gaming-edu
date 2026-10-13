import docx
import re

def replace_text_in_doc(doc, replacements):
    # Iterate through paragraphs in the document
    for p in doc.paragraphs:
        for key, value in replacements.items():
            if key in p.text:
                full_text = p.text.replace(key, value)
                if len(p.runs) > 0:
                    first_run = p.runs[0]
                    bold = first_run.bold
                    italic = first_run.italic
                    underline = first_run.underline
                    font_name = first_run.font.name
                    font_size = first_run.font.size
                    
                    p.clear()
                    new_run = p.add_run(full_text)
                    new_run.bold = bold
                    new_run.italic = italic
                    new_run.underline = underline
                    new_run.font.name = font_name
                    new_run.font.size = font_size
                else:
                    p.text = full_text

    # Iterate through tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for key, value in replacements.items():
                        if key in p.text:
                            full_text = p.text.replace(key, value)
                            if len(p.runs) > 0:
                                first_run = p.runs[0]
                                bold = first_run.bold
                                italic = first_run.italic
                                underline = first_run.underline
                                font_name = first_run.font.name
                                font_size = first_run.font.size
                                
                                p.clear()
                                new_run = p.add_run(full_text)
                                new_run.bold = bold
                                new_run.italic = italic
                                new_run.underline = underline
                                new_run.font.name = font_name
                                new_run.font.size = font_size
                            else:
                                p.text = full_text

base_file = r'c:\Users\LONG\gaming-edu\khao sat game\Bao_cao_tien_do_Tuan_4_Nguyen_Nhat_Long.docx'

weeks = [
    {
        'week': 5,
        'replacements': {
            'TUẦN 4': 'TUẦN 5',
            '24/8/2026 - 30/8/2026': '28/09/2026 - 04/10/2026',
            'Thiết kế UI/UX.': 'Xây dựng giao diện các trang cơ bản (Trang chủ, Đăng nhập, Đăng ký).',
            'Xây dựng cấu trúc dự án phần mềm (Front-end, Back-end).': 'Thiết lập kết nối giữa Front-end và Back-end.',
            'Thiết kế hoàn thiện sơ đồ ERD.': 'Thiết kế UI/UX.',
            'Xây dựng xong cơ sở dữ liệu (Database).': 'Xây dựng cấu trúc dự án phần mềm (Front-end, Back-end).',
            'Quản lý mã nguồn bằng Git/GitHub.': 'Thiết lập môi trường phát triển cục bộ.',
            '25 tháng 09 năm 2026': '29 tháng 09 năm 2026'
        }
    },
    {
        'week': 6,
        'replacements': {
            'TUẦN 4': 'TUẦN 6',
            '24/8/2026 - 30/8/2026': '05/10/2026 - 11/10/2026',
            'Thiết kế UI/UX.': 'Phát triển tính năng quản lý tài khoản người dùng và xác thực.',
            'Xây dựng cấu trúc dự án phần mềm (Front-end, Back-end).': 'Xây dựng chức năng tạo và quản lý bài khảo sát/câu hỏi.',
            'Thiết kế hoàn thiện sơ đồ ERD.': 'Xây dựng giao diện các trang cơ bản (Trang chủ, Đăng nhập, Đăng ký).',
            'Xây dựng xong cơ sở dữ liệu (Database).': 'Thiết lập kết nối giữa Front-end và Back-end.',
            'Quản lý mã nguồn bằng Git/GitHub.': 'Hoàn thiện cấu trúc thư mục của dự án.',
            '25 tháng 09 năm 2026': '06 tháng 10 năm 2026'
        }
    },
    {
        'week': 7,
        'replacements': {
            'TUẦN 4': 'TUẦN 7',
            '24/8/2026 - 30/8/2026': '12/10/2026 - 18/10/2026',
            'Thiết kế UI/UX.': 'Tích hợp tính năng tham gia bài khảo sát theo thời gian thực.',
            'Xây dựng cấu trúc dự án phần mềm (Front-end, Back-end).': 'Phát triển giao diện trình bày tương tác cho người dùng.',
            'Thiết kế hoàn thiện sơ đồ ERD.': 'Phát triển tính năng quản lý tài khoản người dùng và xác thực.',
            'Xây dựng xong cơ sở dữ liệu (Database).': 'Xây dựng chức năng tạo và quản lý bài khảo sát/câu hỏi.',
            'Quản lý mã nguồn bằng Git/GitHub.': 'Tối ưu hóa các API cơ bản cho hệ thống.',
            '25 tháng 09 năm 2026': '13 tháng 10 năm 2026'
        }
    },
    {
        'week': 8,
        'replacements': {
            'TUẦN 4': 'TUẦN 8',
            '24/8/2026 - 30/8/2026': '19/10/2026 - 25/10/2026',
            'Thiết kế UI/UX.': 'Thu thập dữ liệu đánh giá, hiển thị kết quả và biểu đồ.',
            'Xây dựng cấu trúc dự án phần mềm (Front-end, Back-end).': 'Kiểm thử toàn bộ hệ thống và tối ưu hiệu năng.',
            'Thiết kế hoàn thiện sơ đồ ERD.': 'Tích hợp tính năng tham gia bài khảo sát theo thời gian thực.',
            'Xây dựng xong cơ sở dữ liệu (Database).': 'Phát triển giao diện trình bày tương tác cho người dùng.',
            'Quản lý mã nguồn bằng Git/GitHub.': 'Kiểm thử các luồng chức năng đã hoàn thiện.',
            '25 tháng 09 năm 2026': '20 tháng 10 năm 2026'
        }
    }
]

for w in weeks:
    doc = docx.Document(base_file)
    replace_text_in_doc(doc, w['replacements'])
    out_file = fr'c:\Users\LONG\gaming-edu\khao sat game\Bao_cao_tien_do_Tuan_{w["week"]}_Nguyen_Nhat_Long.docx'
    doc.save(out_file)
    print(f'Created {out_file}')

