import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import glob
import os

files = glob.glob(r'c:\Users\LONG\gaming-edu\khao sat game\Bao_cao_tien_do_Tuan_*.docx')
img_path = r'C:\Users\LONG\.gemini\antigravity-ide\brain\46350646-af1e-4d92-acc6-3a4726c6f4c6\.user_uploaded\media_1791542215652.png'

if not os.path.exists(img_path):
    print("Image not found:", img_path)
    exit(1)

for file in files:
    try:
        doc = docx.Document(file)
        
        # In Table 0, Row 0, Cell 1
        table = doc.tables[0]
        cell = table.cell(0, 1)
        
        if "Nguyễn Nhật Long" in cell.text:
            name_idx = -1
            for i, p in enumerate(cell.paragraphs):
                if "Nguyễn Nhật Long" in p.text:
                    name_idx = i
                    break
            
            if name_idx > 0:
                # Clear all paragraphs between the date (0) and name (name_idx)
                for i in range(1, name_idx):
                    p = cell.paragraphs[i]
                    p.clear()
                    for r in p.runs:
                        r.clear()
                
                # We will use paragraph 1 for the signature
                para_to_insert = cell.paragraphs[1]
                para_to_insert.clear()
                
                # Professional formatting for the signature paragraph
                para_to_insert.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para_to_insert.paragraph_format.space_before = Pt(6)
                para_to_insert.paragraph_format.space_after = Pt(6)
                
                run = para_to_insert.add_run()
                # 0.8 inches ensures it doesn't stretch the table too much.
                run.add_picture(img_path, height=Inches(0.8))
                
                doc.save(file)
                print(f"Expertly updated Long's signature with new image in {file}")
            else:
                print(f"Could not find name paragraph in {file}")
        else:
            print(f"Could not find Nguyễn Nhật Long in cell in {file}")
    except Exception as e:
        print(f"Error processing {file}: {e}")
