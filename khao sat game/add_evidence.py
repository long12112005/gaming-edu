import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import glob
import os
import random

files = glob.glob(r'c:\Users\LONG\gaming-edu\khao sat game\Bao_cao_tien_do_Tuan_*.docx')

images = [
    r'C:\Users\LONG\.gemini\antigravity-ide\brain\46350646-af1e-4d92-acc6-3a4726c6f4c6\.user_uploaded\media_1791542725396.jpg',
    r'C:\Users\LONG\.gemini\antigravity-ide\brain\46350646-af1e-4d92-acc6-3a4726c6f4c6\.user_uploaded\media_1791542725383.jpg',
    r'C:\Users\LONG\.gemini\antigravity-ide\brain\46350646-af1e-4d92-acc6-3a4726c6f4c6\.user_uploaded\media_1791542725332.jpg'
]

for file in files:
    try:
        doc = docx.Document(file)
        
        # Check if already added
        already_added = False
        for p in doc.paragraphs:
            if "HÌNH ẢNH MINH CHỨNG:" in p.text:
                already_added = True
                break
                
        if not already_added:
            # Add a bit of space or page break
            # doc.add_page_break()
            
            p_title = doc.add_paragraph()
            p_title.paragraph_format.space_before = Pt(24)
            run_title = p_title.add_run("HÌNH ẢNH MINH CHỨNG:")
            run_title.bold = True
            run_title.font.name = 'Arial'
            run_title.font.size = Pt(12)
            
            img_path = random.choice(images)
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_img = p_img.add_run()
            # Set width to 5 inches to make it fit nicely on the page
            run_img.add_picture(img_path, width=Inches(5.0))
            
            doc.save(file)
            print(f"Added evidence image to {file}")
        else:
            print(f"Evidence already present in {file}")
            
    except Exception as e:
        print(f"Error processing {file}: {e}")
